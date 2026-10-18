"""Word briefing for the RB experiments: conclusions first, one headline + one figure + a few bullets per page.
Tables are computed from the result CSVs in outputs/; the projected-divergence numbers come from
outputs/study_rb_g/unified/summary.txt (analysis/unified_summary.py).
Usage: python analysis/make_report_docx.py --out docs/RB实验汇报.docx
"""
import argparse
from pathlib import Path
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

R = Path('outputs'); G = R / 'study_rb_g'
NAVY, BLUE, LIGHT, GREY = '1F3A5F', '0072B2', 'E8F0F8', 'F4F6F8'
CN, EN = '微软雅黑', 'Calibri'


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color); tcPr.append(sh)


def borders(table, color='BFC7D1', sz=4):
    tblPr = table._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        x = OxmlElement(f'w:{e}'); x.set(qn('w:val'), 'single'); x.set(qn('w:sz'), str(sz)); x.set(qn('w:color'), color); b.append(x)
    tblPr.append(b)


def style_run(run, size=None, bold=None, color=None):
    run.font.name = EN; run._element.rPr.rFonts.set(qn('w:eastAsia'), CN)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = RGBColor.from_string(color)


def field(par, code):
    for t, txt in (('begin', None), (None, code), ('end', None)):
        r = par.add_run(); style_run(r, 9, color='6B7785')
        if t: e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), t); r._r.append(e)
        else: e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt; r._r.append(e)


class Brief:
    def __init__(self):
        d = self.d = Document(); s = d.sections[0]
        s.orientation = WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = Cm(29.7), Cm(21.0)
        s.left_margin = s.right_margin = Cm(2.0); s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.6)
        st = d.styles['Normal']; st.font.name = EN; st.font.size = Pt(11); st.element.rPr.rFonts.set(qn('w:eastAsia'), CN)
        hp = s.header.paragraphs[0]; style_run(hp.add_run('Rayleigh–Bénard 对流：扩散先验稀疏观测同化 · 实验汇报'), 9, color='6B7785')
        fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT; field(fp, 'PAGE')
        self.first = True; self.nfig = 0; self.W = 25.5

    def page(self, title, sub=None):
        if not self.first: self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        self.first = False
        p = self.d.add_paragraph(); style_run(p.add_run(title), 20, True, NAVY); p.paragraph_format.space_after = Pt(2)
        pPr = p._p.get_or_add_pPr(); pb = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
        for k, v in (('val', 'single'), ('sz', '12'), ('space', '1'), ('color', BLUE)): bt.set(qn(f'w:{k}'), v)
        pb.append(bt); pPr.append(pb)
        if sub: q = self.d.add_paragraph(); style_run(q.add_run(sub), 10.5, color='6B7785')

    def headline(self, text):
        t = self.d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; c = t.rows[0].cells[0]
        shade(c, LIGHT); c.text = ''; p = c.paragraphs[0]; p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
        style_run(p.add_run('▌ '), 12, True, BLUE); style_run(p.add_run(text), 12.5, True, NAVY)
        self.d.add_paragraph().paragraph_format.space_after = Pt(0)

    def bullets(self, items, size=11):
        for x in items:
            p = self.d.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(2); style_run(p.add_run(x), size)

    def text(self, x, size=11, color=None, bold=False):
        p = self.d.add_paragraph(); style_run(p.add_run(x), size, bold, color); p.paragraph_format.space_after = Pt(3)

    def table(self, header, rows, widths, best=None, size=10.5, first_left=True):
        t = self.d.add_table(rows=1, cols=len(header)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; borders(t)
        for i, x in enumerate(header):
            c = t.rows[0].cells[i]; shade(c, NAVY); c.text = ''; p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (i == 0 and first_left) else WD_ALIGN_PARAGRAPH.CENTER
            style_run(p.add_run(str(x)), size, True, 'FFFFFF')
        for r, row in enumerate(rows):
            cells = t.add_row().cells
            for i, x in enumerate(row):
                shade(cells[i], GREY if r % 2 else 'FFFFFF'); cells[i].text = ''; p = cells[i].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if (i == 0 and first_left) else WD_ALIGN_PARAGRAPH.CENTER
                style_run(p.add_run(str(x)), size, bold=bool(best and (r, i) in best), color=BLUE if best and (r, i) in best else None)
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
        self.d.add_paragraph().paragraph_format.space_after = Pt(0)

    def fig(self, path, caption, width=None):
        self.nfig += 1
        p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(str(path), width=Cm(width or self.W))
        c = self.d.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_run(c.add_run(f'图 {self.nfig}　{caption}'), 9.5, color='4A5563')


def workflow_fig(path):
    fig, ax = plt.subplots(figsize=(15, 3.1)); ax.axis('off'); ax.set_xlim(0, 15.3); ax.set_ylim(0, 3.1)
    def box(x, y, w, h, t, fc, ec='#1F3A5F', fs=10.5, tc='#1F3A5F'):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle='round,pad=0.05,rounding_size=.12', fc=fc, ec=ec, lw=1.4))
        ax.text(x + w / 2, y + h / 2, t, ha='center', va='center', fontsize=fs, color=tc, linespacing=1.4)
    def arrow(x0, y0, x1, y1): ax.add_patch(FancyArrowPatch((x0, y0), (x1, y1), arrowstyle='-|>', mutation_scale=16, lw=1.6, color='#1F3A5F'))
    box(.2, 1.0, 2.6, 1.1, 'Analysis\n$x^a_{t-1}$', '#E8F0F8')
    box(3.5, 1.0, 3.2, 1.1, 'Dedalus forecast\n(Ra = 5e4, 35 steps)', '#FFF3E0', '#D55E00', tc='#7A3A00')
    box(7.4, 1.0, 2.5, 1.1, 'Background\n$x^b_t$', '#E8F0F8')
    box(10.6, 1.0, 2.3, 1.1, 'Diffusion prior\n+ DPS guidance', '#E3F4EC', '#009E73', tc='#005C42')
    box(13.65, 1.0, 1.3, 1.1, 'Analysis\n$x^a_t$', '#E8F0F8')
    box(7.4, 2.45, 5.8, .55, 'Observations $y_t$ : 1% of grid points, 3 channels, noise-free', '#F4F6F8', '#6B7785', 9.5, '#333')
    arrow(2.9, 1.55, 3.4, 1.55); arrow(6.8, 1.55, 7.3, 1.55); arrow(10.0, 1.55, 10.5, 1.55); arrow(13.0, 1.55, 13.55, 1.55)
    arrow(11.75, 2.4, 11.75, 2.2)
    ax.text(7.5, .35, 'Background and observations both enter DPS as likelihood guidance (sampling starts from pure noise);  the cycle repeats at every frame.', ha='center', fontsize=10, color='#555')
    fig.savefig(path, dpi=200, bbox_inches='tight'); plt.close(fig)


def f4(x): return f'{x:.4f}'


def main():
    ap = argparse.ArgumentParser(description=__doc__); ap.add_argument('--out', default='docs/RB实验汇报.docx'); a = ap.parse_args()
    s = []
    for c in range(3):
        u = pd.read_csv(G / f'unified/c{c}/enkf.csv'); f = pd.read_csv(G / f'fusion_c{c}/fusion.csv')
        s += [u[u.scheme == 'dps'].assign(case=c, arm='A'), f[f.scheme.str.contains('only')].assign(case=c, arm='B'), u[u.scheme == 'enkf_infl_loc'].assign(case=c, arm='C')]
    s = pd.concat(s)
    cols = ['missing_rmse_ch0', 'missing_rmse_ch1', 'missing_rmse_ch2', 'missing_rmse']
    ch = s.groupby('arm')[cols].mean(); pc = s.groupby(['arm', 'case']).missing_rmse.mean().unstack()
    div = s.groupby('arm').an_div_rel.median(); tdiv = s.truth_div_rel.median()
    A, B, C = (ch.loc[k, 'missing_rmse'] for k in 'ABC')
    e3m = pd.read_csv(G / 'enkf3/enkf_all.csv').groupby('scheme').missing_rmse.mean()
    noise = pd.read_csv(G / 'noise/noise.csv'); nz = noise[noise.level.isin([0.0, 0.4])]
    mm = nz.groupby(['scheme', 'level']).missing_rmse.mean(); lt = nz[nz.frame >= 90].groupby(['scheme', 'level']).missing_rmse.mean()
    long = pd.concat([pd.read_csv(f) for f in sorted((R / 'study_rb_long').glob('c*/enkf.csv'))])
    lw = long.assign(w=long.frame // 50 * 50).groupby(['scheme', 'w']).missing_rmse.mean().unstack(0)
    la = long[long.scheme == 'dps'].missing_rmse.mean(); le = long[long.scheme == 'enkf_infl_loc'].missing_rmse.mean()
    wf = G / 'unified' / 'workflow.png'; workflow_fig(wf)
    NP, NA = 'No assimilation (forecast only)', 'DPS, observations every frame'
    D = Brief()

    # ---- 1 cover + key conclusions
    t = D.d.add_paragraph(); style_run(t.add_run('扩散先验稀疏观测同化'), 30, True, NAVY); t.paragraph_format.space_after = Pt(0)
    t = D.d.add_paragraph(); style_run(t.add_run('Rayleigh–Bénard 对流 · 实验汇报'), 18, False, BLUE)
    D.text('1% 稀疏观测 · 数值模式带模型误差 · 逐帧循环同化 · 孪生实验', 11, '6B7785')
    D.headline('核心结论')
    D.bullets([f'精度：背景场 + 稀疏观测联合同化（A）未观测点 RMSE 最低，为 {f4(A)}；较稀疏观测直接重建（B）降低 {100 * (1 - A / B):.0f}%，较 EnKF（C）降低 {100 * (1 - A / C):.0f}%。',
               f'稳定性：A 相对 B 的优势在 3 个测试 case 中 2 个显著、1 个持平；循环初期（约前 6 个时间单位）背景场反而不利。',
               f'平衡：原始分析场散度 A 约 {div["A"]:.2f}、C 约 {div["C"]:.3f}（真值 {tdiv:.0e}）；经求解器同款流函数投影后三者均约 0.008，无明显差异，且速度误差基本不变。',
               f'初值噪声：40% 初值噪声使纯预报误差增大约 {100 * (mm[NP, 0.4] / mm[NP, 0.0] - 1):.0f}%，同化后几乎不受影响。',
               f'长时间：400 个循环内 A 稳定在约 0.011，EnKF 约 0.019。'], 12)
    D.d.add_paragraph()
    D.table(['', 'A　背景场 + 观测（DPS）', 'B　仅观测（DPS）', 'C　EnKF'], [
        ['未观测点 RMSE', f4(A), f4(B), f4(C)], ['相对 A', '—', f'+{100 * (B / A - 1):.0f}%', f'+{100 * (C / A - 1):.0f}%']],
        [5.5, 6.5, 6.5, 6.5], best={(0, 1)}, size=12)
    D.text('RMSE：三通道合并、仅统计未观测格点；3 个测试 case × 100 帧平均；每帧进观测。', 9.5, '6B7785')

    # ---- 2 setup
    D.page('实验设置', '孪生实验：真值由 Dedalus 在 Ra = 1×10⁵ 下生成，观测取自真值')
    D.table(['项目', '设置'], [
        ['控制方程', '二维 Boussinesq Rayleigh–Bénard：∇·u = 0；∂ₜb − κ∇²b = −u·∇b；∂ₜu − ν∇²u + ∇p − b e_z = −u·∇u；κ = (Ra·Pr)^(−1/2)，ν = (Ra/Pr)^(−1/2)'],
        ['边界与网格', 'x 周期（L_x = 3），z 上下壁面（L_z = 1）：b(0) = 1、b(L_z) = 0、壁面无滑移；Fourier × Chebyshev，72 × 24，dt = 0.01'],
        ['状态与数据', '状态 (b, u_x, u_z)；200 条轨迹，弃前 120 帧瞬态，每条 100 帧；训练 / 验证 / 测试 = 160 / 20 / 20；测试用 case 0–2'],
        ['时间尺度', '1 帧 = 35 个求解器步 = 0.35 时间单位；100 帧 = 3500 步 = 35 时间单位'],
        ['模型误差', '背景模型 Ra = 5×10⁴（真值 1×10⁵）参数失配；误差由求解器自身产生，不人为加噪'],
        ['观测', '1%（17 / 1728 个格点），位置固定，三通道同时观测，无噪声；每帧进一次观测'],
        ['先验', '无条件扩散模型（VP-cosine，U-Net，c = 16），只学习流场快照分布，200 epoch'],
        ['后验采样', 'DPS，100 步、8 样本取均值；关闭硬观测覆盖；参数在验证 case 上选定'],
        ['评价指标', '未观测点 RMSE（三通道合并）；相对散度 ‖∇·u‖/‖∇u‖（Fourier–Chebyshev 谱导数）']],
        [4.2, 21.3], size=10.5)
    D.fig(wf, '逐帧循环同化流程（方案 A）', width=22)

    # ---- 3 schemes
    D.page('对照方案', '三个方案在同一观测、同一预报模型、同一硬观测策略下比较')
    D.table(['方案', '方法', '背景场', '主要参数（均在验证 case 上调定）'], [
        ['A', '背景场 + 稀疏观测联合同化', 'Dedalus 预报（Ra = 5e4）', 'DPS：bg_std 0.3、dps_rate 2、obs_std 0.05'],
        ['B', '稀疏观测直接重建', '无（每帧独立采样）', 'DPS：obs_std 0.025、dps_rate 2'],
        ['C', '随机 EnKF', '成员由 Dedalus 预报', '20 成员；膨胀 2.5；Gaspari–Cohn 局地化 1.5(x) × 1.2(z)；R = 0.01']],
        [2, 6.5, 6.5, 10.5], size=11)
    D.bullets(['EnKF 更新式：xᵃ − xᶠ = Pᶠ Hᵀ (H Pᶠ Hᵀ + R)⁻¹ [y − H(xᶠ)]，扰动观测、逐观测点顺序同化；分析场取集合均值。',
               'EnKF 初始集合以 DPS 第 0 帧分析场为中心，故比较的是同一初始分析之后的循环。',
               'EnKF 散度在预报重启之前测量；DPS 与 EnKF 均不做硬观测覆盖。'], 10.5)

    # ---- 4 main result
    D.page('结果 1　重建误差', '每帧进观测，3 个测试 case，100 帧')
    D.headline(f'A 的未观测点 RMSE 最低：{f4(A)}，较 B 降低 {100 * (1 - A / B):.0f}%，较 EnKF 降低 {100 * (1 - A / C):.0f}%')
    rows = [[k + '　' + n] + [ch.loc[k, c] for c in cols] + [pc.loc[k, c] for c in range(3)] for k, n in (('A', '背景场 + 观测'), ('B', '仅观测'), ('C', 'EnKF'))]
    best = {(min(range(3), key=lambda r: rows[r][j]), j) for j in range(1, 8)}
    D.table(['方案', 'b', 'u_x', 'u_z', '合并', 'case 0', 'case 1', 'case 2'], [[r[0]] + [f4(x) for x in r[1:]] for r in rows],
            [5.2, 2.6, 2.6, 2.6, 2.6, 3.3, 3.3, 3.3], best=best)
    D.fig(G / 'unified/unified.png', '(a) 分通道平均 RMSE；(b) RMSE 随时间；(c) 分析场相对散度（虚线：真值）')
    D.bullets([f'A 在 case 1、2 显著领先；case 0 与 B 持平（{f4(pc.loc["A", 0])} 对 {f4(pc.loc["B", 0])}）。',
               f'EnKF 的误差集中在 b 通道（{f4(ch.loc["C", cols[0]])}），u_x、u_z 与 B 接近。',
               f'每 3 帧进观测时：A {f4(e3m["dps"])}，EnKF（膨胀 + 局地化）{f4(e3m["enkf_infl_loc"])}，仅膨胀 {f4(e3m["enkf_inflation"])}，无膨胀无局地化 {f4(e3m["enkf_plain"])}（几乎不吸收观测）。'], 10.5)

    # ---- 5 background
    D.page('结果 2　背景场的作用', 'A 与 B 的逐帧 RMSE')
    D.headline('背景场平均有益，但不稳定：2 / 3 个 case 显著受益，初期反而不利')
    D.fig(G / 'fusion/fusion.png', 'A（蓝）与 B（橙）的 RMSE 随时间；标题给出各 case 的平均值')
    D.bullets(['循环初期（约前 6 个时间单位）：上一步分析尚不准，预报出的背景场拖累结果，A 不如 B。',
               'B 每帧独立采样、没有时间连续性；A 的优势中有一部分来自求解器带来的时间连续性，而非背景场精度本身。'], 10.5)

    # ---- 6 fields
    D.page('结果 3　流场对比', '测试 case 1，第 60 帧，浮力 b')
    D.fig(G / 'unified/field_b_c1_preview.png', '真值、观测位置，A / B / C 的分析场与绝对误差（误差色标统一）；底部为 RMSE 随时间曲线', width=15.5)
    D.text('动画（3 个 case × 3 个变量）：outputs/study_rb_g/gif/', 10, '6B7785')

    # ---- 7 divergence
    D.page('结果 4　方程平衡（速度散度）', '相对散度 ‖∇·u‖/‖∇u‖，谱导数；真值约 3.4×10⁻⁵')
    D.headline('经求解器同款流函数投影后，三方案散度无明显差异；原始分析场上扩散模型更大')
    D.table(['方案', '原始分析场', '投影后', '速度 RMSE：投影前', '投影后'], [
        ['A　背景场 + 观测', f'{div["A"]:.3f}', '0.0084', '0.01348', '0.01335'],
        ['B　仅观测', f'{div["B"]:.3f}', '0.0084', '0.01744', '0.01690'],
        ['C　EnKF', f'{div["C"]:.4f}', '0.0075', '0.01562', '0.01564'],
        ['参照：真值经同样投影', f'{tdiv:.1e}', '0.0079', '—', '—']], [7, 4.5, 3.5, 5.5, 4.5], size=11)
    D.bullets(['投影即每帧预报前求解器重启所做的操作，因此原始分析场的不平衡不会累积。',
               '投影后三者均贴近投影算子下限（0.0079），且速度误差基本不变：原始散度偏大主要是小尺度毛刺被求导放大（给真值速度加 0.6% 白噪声即可使散度升至 0.06）。该解释在 3 个 case、102 帧上验证。',
               '无局地化 EnKF 的散度约 1×10⁻⁵ 并非同化质量所致：成员均由求解器预报、近似无散度，其线性组合亦然（纯预报约 9×10⁻⁶）；局地化破坏该性质，升至 10⁻³ 量级。'], 10.5)

    # ---- 8 noise
    D.page('结果 5　初值噪声', 't = 0 叠加平滑噪声，水平 0–40% 参考标准差，测试 case 0，3 个噪声样本')
    D.headline(f'40% 初值噪声使纯预报误差增大 {100 * (mm[NP, 0.4] / mm[NP, 0.0] - 1):.0f}%，每帧同化后几乎不变')
    D.fig(G / 'noise/noise.png', '上排：未观测点 RMSE（橙：纯预报，蓝：DPS）；下排：同化后相对散度', width=22)
    D.table(['方案', '无噪声', '40% 噪声', '变化', '末 10 帧：无噪声', '40% 噪声'], [
        [n, f4(mm[k, 0.0]), f4(mm[k, 0.4]), f'{100 * (mm[k, 0.4] / mm[k, 0.0] - 1):+.0f}%', f4(lt[k, 0.0]), f4(lt[k, 0.4])]
        for k, n in ((NP, '纯预报'), (NA, 'DPS 每帧观测'), ('DPS, observations every 3 frames', 'DPS 每 3 帧观测'))],
        [5, 3.4, 3.4, 3, 4.5, 3.4], size=10)
    D.bullets(['噪声构造：b 为平滑噪声乘壁面包络；速度由流函数噪声求导得到（无散度、壁面为零），u_x 与 u_z 共用一个缩放；幅度 = 水平 × 参考标准差（b：0.097；u_x、u_z：0.166、0.169）。',
               '末 10 帧纯预报已因模型误差涨至约 0.052，初值噪声影响被淹没，故该实验区分度偏弱。'], 10)

    # ---- 9 long
    D.page('结果 6　长时间表现', '3 条轨迹重积分至 520 帧，取 400 个循环（约 140 时间单位），每 3 帧进观测')
    D.headline(f'DPS 约 100 帧后收敛至约 0.011；EnKF 全程约 0.019（全程平均 {f4(la)} 对 {f4(le)}）')
    D.fig(R / 'study_rb_long/long.png', '(a) 400 个循环的 RMSE；(b) 每 17.5 时间单位窗口的平均 RMSE', width=21)
    D.bullets(['真值在此区间已是统计稳态（动能约 0.0276，b 标准差约 0.257，帧间变化约 0.0127，均不随时间变化），流场仍在演化。',
               '误差走平是否源于"模型误差注入与观测修正的平衡"尚未验证（验证办法：改变 Ra 失配量，看平线高度是否随之变化）。'], 10.5)

    # ---- 10 limits
    D.page('局限与后续')
    D.bullets(['样本：3 个测试 case、单个随机种子；初值噪声实验仅用测试 case 0。',
               '设定：观测无噪声、三通道同时观测、模型误差仅 Ra 失配一种，较真实场景容易。',
               'EnKF：仅 20 个成员（未测 40 及以上）；更新式按通用形式实现，待与任务给定公式核对。',
               '散度：原始散度偏大的解释仅在 102 帧样本上验证。',
               '后续：多种子与更多 case；EnKF 增加成员；改变 Ra 失配量与观测比例；观测加噪。'], 12)
    D.text('主要文件：outputs/study_rb_g/{unified, fusion_c0..2, noise, gif}、outputs/study_rb_long；脚本 analysis/{study_rb, plot_study_rb, make_rb_gif, unified_summary}.py；详见 README §6。', 9.5, '6B7785')
    out = Path(a.out); out.parent.mkdir(parents=True, exist_ok=True); D.d.save(out); print('saved', out)


if __name__ == '__main__':
    main()
